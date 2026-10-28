"""Generate the English Task 2 Chat Backend FUT from actual pytest JUnit evidence.

Run the existing seven chat/LLM/prompt test modules with --junitxml=tmp/chat-fut.xml.
Then run: python generate_chat_fut.py --junit tmp/chat-fut.xml

The style helpers come from the existing sibling generate_auth_fut.py.
No APIs, migrations or application tests run when rendering this workbook.
Without execution evidence, all cases remain Not Run.
"""
from __future__ import annotations

import argparse
import hashlib
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

from generate_auth_fut import (
    BORDER, HEADERS, ORANGE, RESULT_COLORS, TESTER, WIDTHS,
    estimated_height, font, merged,
)

OUTPUT = Path(__file__).resolve().with_name("FUT_BE_Chat_Conversation_Management.xlsx")
SOURCE_REVISION = "3f4e13b"
MODULE = "Complete the Chat Backend and Account-Based Conversation Management"
MANUAL_QA = (
    "Cao Văn Khoa confirmed that all manual Task 2 acceptance checks passed. "
    "This is tester-reported sign-off; automated case results below come only from the supplied JUnit report."
)
TEST_CLASSES = set([
    "apps.backend.tests.test_chat_conversation.ChatConversationTests",
    "apps.backend.tests.test_chat_api.ChatApiTests",
    "apps.backend.tests.test_chat_fallback.DevelopmentAnswerTests",
    "apps.backend.tests.test_chat_fallback.LocalConversationFallbackTests",
    "apps.backend.tests.test_chat_integration.ChatPipelineIntegrationTests",
    "apps.backend.tests.test_chat_integration.ChatStartupTests",
    "apps.backend.tests.test_chat_graphrag.ChatGraphRagTests",
    "apps.backend.tests.test_llama_server.LlamaServerTests",
    "apps.backend.tests.test_prompt.OverviewPromptTests"
])
CASES = [
    {
        "id": "CH_CON_01",
        "group": "1. Chat Request and Response Contracts",
        "description": "Create a persisted chat response with normalized citation fields",
        "procedure": "1. Authenticate an active account.\n2. POST /api/chat/message with a valid question and conversation_id=null.\n3. Inspect the response and citation linkage.",
        "expected": "HTTP 201; conversation_id, message_id, answer and normalized citations are returned. Citations link to the assistant message.",
        "actual": "HTTP 201; normalized source title, URL, relevance score and assistant message linkage were verified.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_create_message_persists_conversation_both_turns_and_citations"
        ]
    },
    {
        "id": "CH_CON_02",
        "group": "1. Chat Request and Response Contracts",
        "description": "Keep the compatibility chat response available",
        "procedure": "1. POST /api/chat with a valid question.\n2. Provide controlled retrieved evidence and a model answer.\n3. Inspect the compatibility response.",
        "expected": "HTTP 200; answer and sources are returned, with answer_type=answered and resolution_status=resolved.",
        "actual": "HTTP 200; answer, passage source, answer_type and resolution_status matched the compatibility contract.",
        "tests": [
            "apps.backend.tests.test_chat_api.ChatApiTests::test_returns_answer_with_passage_citation"
        ]
    },
    {
        "id": "CH_CON_03",
        "group": "1. Chat Request and Response Contracts",
        "description": "Reject invalid input and client-supplied account identity",
        "procedure": "1. POST /api/chat/message with blank/oversized messages, a malformed conversation UUID, account_id in the body or oversized heritage/scene context.\n2. Inspect validation and inference calls.",
        "expected": "HTTP 422 for each invalid payload; no AI retrieval or account/conversation writes occur.",
        "actual": "All invalid payloads returned HTTP 422 before retrieval; no conversation, message or citation was stored.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_invalid_requests_and_account_spoofing_are_rejected_before_ai"
        ]
    },
    {
        "id": "CH_CON_04",
        "group": "1. Chat Request and Response Contracts",
        "description": "Preserve the compatibility endpoint blank-message error",
        "procedure": "1. POST /api/chat with a whitespace-only message.\n2. Inspect the HTTP status.",
        "expected": "HTTP 400; a blank question is rejected.",
        "actual": "The compatibility endpoint rejected the blank question with HTTP 400.",
        "tests": [
            "apps.backend.tests.test_chat_api.ChatApiTests::test_blank_message_is_rejected"
        ]
    },
    {
        "id": "CH_CON_05",
        "group": "1. Chat Request and Response Contracts",
        "description": "Accept valid maximum message and UI context sizes",
        "procedure": "1. POST /api/chat/message with an 8000-character message and 2000-character heritage and scene context fields.\n2. Use valid controlled evidence.\n3. Inspect the answer and retrieval input.",
        "expected": "The valid request succeeds; the fields are not incorrectly revalidated as one oversized message.",
        "actual": "The request succeeded with answer_type=answered; the maximum heritage context reached retrieval.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_valid_maximum_message_and_ui_context_are_not_revalidated_as_one_oversized_message"
        ]
    },
    {
        "id": "CH_CON_06",
        "group": "1. Chat Request and Response Contracts",
        "description": "Apply context validation to the compatibility interface",
        "procedure": "1. POST /api/chat with a developer-role context message.\n2. Repeat with seven context messages.\n3. Inspect both responses.",
        "expected": "Both requests return HTTP 422.",
        "actual": "Both disallowed-role and excessive-context requests returned HTTP 422.",
        "tests": [
            "apps.backend.tests.test_chat_api.ChatApiTests::test_legacy_context_uses_the_same_role_and_size_guards"
        ]
    },
    {
        "id": "CH_AI_01",
        "group": "2. AI Pipeline, Citations, Clarification and Abstention",
        "description": "Return an answer grounded in retrieved passage evidence",
        "procedure": "1. Supply a known passage and a controlled model response to POST /api/chat.\n2. Inspect the answer and passage citation.",
        "expected": "The answer retains the retrieved passage citation and is classified as answered/resolved.",
        "actual": "The returned answer and passage citation matched the controlled evidence; the outcome was answered/resolved.",
        "tests": [
            "apps.backend.tests.test_chat_api.ChatApiTests::test_returns_answer_with_passage_citation"
        ]
    },
    {
        "id": "CH_AI_02",
        "group": "2. AI Pipeline, Citations, Clarification and Abstention",
        "description": "Return and restore clarification requests",
        "procedure": "1. Simulate an ambiguous heritage query with candidate suggestions.\n2. Send a persisted chat request.\n3. Restore its assistant message through the history API.",
        "expected": "needs_user_choice=true; suggestions retained; abstained=false; citations may be empty. Clarification metadata survives restoration without generation.",
        "actual": "Clarification flags and suggestions were returned and restored unchanged; no model generation was invoked.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_clarification_metadata_survives_persistence_and_history_restore"
        ]
    },
    {
        "id": "CH_AI_03",
        "group": "2. AI Pipeline, Citations, Clarification and Abstention",
        "description": "Return and restore an evidence-based abstention",
        "procedure": "1. Make retrieval return no evidence.\n2. Send a persisted chat request.\n3. Retrieve the message history.",
        "expected": "abstained=true; abstention_reason=insufficient_evidence; no invented citations. The same outcome appears in history.",
        "actual": "The abstention flag and reason survived restoration; citations remained empty.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_abstention_metadata_survives_history_restore_without_invented_citations"
        ]
    },
    {
        "id": "CH_AI_04",
        "group": "2. AI Pipeline, Citations, Clarification and Abstention",
        "description": "Restore corrections and answer metadata",
        "procedure": "1. Provide retrieval corrections for an unaccented heritage name.\n2. Send a persisted message.\n3. Read back the assistant message.",
        "expected": "corrected_from, corrected_to, intent and abstention state remain available after restoration.",
        "actual": "Corrections, intent and the non-abstention state matched the original response.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_answer_metadata_and_corrections_are_restored"
        ]
    },
    {
        "id": "CH_AI_05",
        "group": "2. AI Pipeline, Citations, Clarification and Abstention",
        "description": "Use original passages rather than generated GraphRAG claims as evidence",
        "procedure": "1. Supply GraphRAG lineage pointing to an original passage.\n2. Include an unsupported generated claim in the graph response.\n3. Inspect the evidence passed into the chat answer.",
        "expected": "Original passage facts and valid graph context reach generation; unsupported generated claims do not. Passage provenance is retained.",
        "actual": "Original passage facts and context were preserved; the unsupported generated claim was excluded and passage provenance was retained.",
        "tests": [
            "apps.backend.tests.test_chat_graphrag.ChatGraphRagTests::test_original_chunk_not_generated_graph_answer_reaches_llm"
        ]
    },
    {
        "id": "CH_AI_06",
        "group": "2. AI Pipeline, Citations, Clarification and Abstention",
        "description": "Reject stale or incomplete source lineage",
        "procedure": "1. Simulate mismatched content hashes, object keys and buckets.\n2. Simulate graph context with incomplete passage lineage.\n3. Inspect accepted evidence and graph context.",
        "expected": "Stale evidence and unsupported graph context are excluded from the prompt.",
        "actual": "Stale mappings produced no accepted evidence; incomplete graph lineage was excluded from graph context.",
        "tests": [
            "apps.backend.tests.test_chat_graphrag.ChatGraphRagTests::test_stale_mapping_or_incomplete_lineage_never_reaches_prompt"
        ]
    },
    {
        "id": "CH_AI_07",
        "group": "2. AI Pipeline, Citations, Clarification and Abstention",
        "description": "Preserve hybrid evidence when optional GraphRAG fails",
        "procedure": "1. Provide valid hybrid retrieval evidence.\n2. Make optional GraphRAG raise a timeout.\n3. Inspect the final retrieval evidence.",
        "expected": "Hybrid source evidence remains available; failed graph context is omitted.",
        "actual": "Hybrid evidence was preserved while graph context was empty.",
        "tests": [
            "apps.backend.tests.test_chat_graphrag.ChatGraphRagTests::test_graph_failure_preserves_hybrid_evidence"
        ]
    },
    {
        "id": "CH_AI_08",
        "group": "2. AI Pipeline, Citations, Clarification and Abstention",
        "description": "Use the existing llama.cpp HTTP inference contract",
        "procedure": "1. Invoke the real llama.cpp HTTP adapter with a controlled HTTP response.\n2. Inspect the outgoing request and parsed generation metadata.",
        "expected": "POST /v1/chat/completions uses the message payload, configured token budget, temperature and timeout; model and token usage are captured.",
        "actual": "The OpenAI-compatible request payload and timeout matched the adapter contract; model and usage were parsed.",
        "tests": [
            "apps.backend.tests.test_llama_server.LlamaServerTests::test_generate_posts_openai_compatible_chat_request"
        ]
    },
    {
        "id": "CH_CTX_01",
        "group": "3. Heritage, Scene and Bounded Conversation Context",
        "description": "Qualify the question with heritage and scene context",
        "procedure": "1. Send heritage_id, heritage_context and scene_context with a persisted question.\n2. Inspect retrieval and generation inputs.",
        "expected": "Heritage/scene context qualifies retrieval; only retrieved evidence is used as the generation source. The heritage association is retained.",
        "actual": "Heritage and scene appeared in retrieval, generation used the passage evidence, and the heritage ID was restored.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_heritage_and_scene_context_qualify_query_without_becoming_source_evidence"
        ]
    },
    {
        "id": "CH_CTX_02",
        "group": "3. Heritage, Scene and Bounded Conversation Context",
        "description": "Resolve heritage context from the stored heritage association",
        "procedure": "1. Create a conversation associated with a published heritage ID.\n2. Omit explicit heritage_context.\n3. Continue with a follow-up question.",
        "expected": "The server loads the published heritage identity and uses it to qualify both requests.",
        "actual": "The published heritage title qualified the original question and follow-up without client-supplied heritage context.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_heritage_identity_from_server_qualifies_followup_without_client_context"
        ]
    },
    {
        "id": "CH_CTX_03",
        "group": "3. Heritage, Scene and Bounded Conversation Context",
        "description": "Load owned server history for retrieval and generation",
        "procedure": "1. Create an account-owned conversation.\n2. Send a contextual follow-up using its ID.\n3. Inspect history, retrieval and generation inputs.",
        "expected": "Prior user/assistant turns reach the model. Retrieval may use the prior user question, while prior assistant text is not treated as evidence.",
        "actual": "Owned turns reached generation; retrieval included the previous user question and excluded the previous assistant claim.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_owned_history_is_loaded_automatically_and_reaches_retrieval_and_generation"
        ]
    },
    {
        "id": "CH_CTX_04",
        "group": "3. Heritage, Scene and Bounded Conversation Context",
        "description": "Resolve a short referential follow-up",
        "procedure": "1. Ask a heritage-specific question.\n2. Continue the same conversation with a short follow-up such as a construction-date question.\n3. Inspect retrieval.",
        "expected": "Retrieval retains the subject of the previous user question.",
        "actual": "The previous heritage subject was present in the follow-up retrieval query.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_short_followup_uses_the_previous_site_question"
        ]
    },
    {
        "id": "CH_CTX_05",
        "group": "3. Heritage, Scene and Bounded Conversation Context",
        "description": "Keep explicit topic changes separate during retrieval",
        "procedure": "1. Start a conversation about one heritage.\n2. Ask explicitly about a different heritage.\n3. Inspect the retrieval query and retained model history.",
        "expected": "The new retrieval query does not mix in the previous heritage; bounded model history can still be retained.",
        "actual": "The old heritage was absent from the new retrieval query; the explicit new subject and prior turns were handled separately.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_explicit_topic_switch_does_not_mix_previous_question_into_retrieval"
        ]
    },
    {
        "id": "CH_CTX_06",
        "group": "3. Heritage, Scene and Bounded Conversation Context",
        "description": "Bound persisted context size before inference",
        "procedure": "1. Create a conversation containing long turns.\n2. Send additional follow-ups.\n3. Inspect the history forwarded to generation.",
        "expected": "At most six messages and 6000 total content characters, with at most 2000 per message, are forwarded. The selected history starts with a user turn.",
        "actual": "Forwarded history met all three limits and began with a user message.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_large_persisted_turns_are_clipped_to_the_context_budget"
        ]
    },
    {
        "id": "CH_CTX_07",
        "group": "3. Heritage, Scene and Bounded Conversation Context",
        "description": "Reject untrusted roles, oversized context and extra history fields",
        "procedure": "1. Submit system-role context, extra account_id fields, a 2001-character context message, seven messages and a total exceeding 6000 characters.\n2. Inspect responses and stored rows.",
        "expected": "HTTP 422 for every invalid context; no conversation data stored.",
        "actual": "All invalid context variants returned HTTP 422 and the database remained unchanged.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_untrusted_history_roles_sizes_and_extra_fields_are_rejected"
        ]
    },
    {
        "id": "CH_CTX_08",
        "group": "3. Heritage, Scene and Bounded Conversation Context",
        "description": "Reject client replacement of an owned conversation history",
        "procedure": "1. Create an owned conversation.\n2. Append a message with that ID and a nonempty forged conversation_context.\n3. Inspect the response and existing history.",
        "expected": "HTTP 400; the original stored messages and citations remain unchanged.",
        "actual": "The replacement attempt returned HTTP 400; the existing conversation, message pair and citation were preserved.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_client_history_cannot_replace_owned_server_history"
        ]
    },
    {
        "id": "CH_CTX_09",
        "group": "3. Heritage, Scene and Bounded Conversation Context",
        "description": "Accept bounded guest context without importing fabricated history",
        "procedure": "1. Send a guest question with bounded user/assistant client context and no saved conversation ID.\n2. Inspect generation inputs and persisted rows.",
        "expected": "Client context reaches generation; only the new message pair is stored. Client context does not grant saved-history access.",
        "actual": "Guest context reached generation, only one new pair was stored, and an account could not read the guest history.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_guest_can_supply_bounded_context_without_claiming_saved_conversation"
        ]
    },
    {
        "id": "CH_CTX_10",
        "group": "3. Heritage, Scene and Bounded Conversation Context",
        "description": "Forward conversation turns to both inference adapters",
        "procedure": "1. Supply bounded user/assistant history to llama.cpp generation.\n2. Verify Gemini request role mapping using its controlled HTTP transport.\n3. Inspect the current source evidence placement.",
        "expected": "History is separate from source evidence. llama.cpp keeps user/assistant roles; Gemini maps assistant to model and appends the current user/evidence request.",
        "actual": "Both adapters transmitted prior turns in the expected roles while preserving the current source evidence.",
        "tests": [
            "apps.backend.tests.test_llama_server.LlamaServerTests::test_generation_transmits_bounded_history_to_the_model_request",
            "apps.backend.tests.test_chat_fallback.DevelopmentAnswerTests::test_gemini_http_contract_uses_key_header_and_rejects_incomplete_output"
        ]
    },
    {
        "id": "CH_PER_01",
        "group": "4. Account Ownership and Persistent Storage",
        "description": "Persist the conversation, both turns and normalized citations",
        "procedure": "1. Send an authenticated new chat message.\n2. Inspect the conversation, user/assistant messages and citation records.",
        "expected": "The verified account owns the conversation; one conversation, two turns and the expected citation are stored atomically.",
        "actual": "The stored owner matched the verified account; conversation, message pair, citation linkage and original passage provenance were verified.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_create_message_persists_conversation_both_turns_and_citations"
        ]
    },
    {
        "id": "CH_PER_02",
        "group": "4. Account Ownership and Persistent Storage",
        "description": "Continue the same conversation and preserve turn order",
        "procedure": "1. Create a conversation.\n2. Append another question with the same conversation ID.\n3. Read its history.",
        "expected": "The conversation ID stays unchanged; four ordered turns are stored as user/assistant/user/assistant.",
        "actual": "The ID remained unchanged and the four turns were restored in the expected order.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_continuation_keeps_one_conversation_and_orders_turns"
        ]
    },
    {
        "id": "CH_PER_03",
        "group": "4. Account Ownership and Persistent Storage",
        "description": "Preserve the response citation order in history",
        "procedure": "1. Return two passages with citation UUIDs intentionally ordered differently from retrieval order.\n2. Persist the answer.\n3. Restore the message.",
        "expected": "The restored citation list exactly matches the response citation order.",
        "actual": "The restored citation list matched the response despite the opposing UUID order.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_multiple_citations_keep_response_order_when_history_is_restored"
        ]
    },
    {
        "id": "CH_PER_04",
        "group": "4. Account Ownership and Persistent Storage",
        "description": "Enforce the database evidence gate",
        "procedure": "1. In the isolated database, try to insert a factual assistant message with no citations.\n2. Leave both abstention and clarification flags false.\n3. Flush the transaction.",
        "expected": "The database rejects the unsupported factual answer with an integrity error.",
        "actual": "The evidence constraint rejected the unsupported assistant message and the transaction was rolled back.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_uncited_nonabstaining_nonclarification_answers_violate_database_gate"
        ]
    },
    {
        "id": "CH_PER_05",
        "group": "4. Account Ownership and Persistent Storage",
        "description": "Reject a message without a valid parent",
        "procedure": "1. In the isolated database, attempt to store a chat message without a valid conversation/session parent.\n2. Flush the transaction.",
        "expected": "The database rejects the invalid message-parent relationship.",
        "actual": "The invalid message was rejected by the database integrity constraints.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_database_rejects_a_message_without_a_parent"
        ]
    },
    {
        "id": "CH_PER_06",
        "group": "4. Account Ownership and Persistent Storage",
        "description": "Cascade owner deletion and preserve unrelated legacy data",
        "procedure": "1. Create an owned conversation with messages, citations and feedback.\n2. DELETE /api/chat/conversations/{id} as its owner.\n3. Check cascades and an unrelated legacy session message.",
        "expected": "HTTP 204 with an empty body; associated messages/citations/feedback are removed. The unrelated legacy session data remains.",
        "actual": "Owner deletion returned an empty HTTP 204; dependent rows were removed and unrelated legacy session data was preserved.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_owner_delete_returns_empty_204_and_cascades_messages_citations_feedback",
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_deleting_conversation_does_not_remove_legacy_session_messages"
        ]
    },
    {
        "id": "CH_HIS_01",
        "group": "5. Conversation Listing and Message History APIs",
        "description": "List only the signed-in account conversations",
        "procedure": "1. Create conversations for account A, account B and a guest.\n2. GET /api/chat/conversations as A.\n3. Check whether listing opens an AI dependency.",
        "expected": "HTTP 200; only A conversations are listed. Listing does not require retrieval/inference.",
        "actual": "Only the owned conversation was listed and no retrieval dependency was opened.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_list_contains_only_authenticated_accounts_own_conversations"
        ]
    },
    {
        "id": "CH_HIS_02",
        "group": "5. Conversation Listing and Message History APIs",
        "description": "Return ordered history and restored citations",
        "procedure": "1. Create a conversation with a complete answer.\n2. GET /api/chat/conversations/{id}/messages as the owner.",
        "expected": "HTTP 200; user then assistant order; no user citations; assistant citations match the original response.",
        "actual": "The history endpoint returned the ordered turns and matching normalized assistant citations.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_history_returns_ordered_messages_and_normalized_citations"
        ]
    },
    {
        "id": "CH_HIS_03",
        "group": "5. Conversation Listing and Message History APIs",
        "description": "Apply bounded pagination to lists and history",
        "procedure": "1. Create several conversations.\n2. Query list/history with limit=1 and offset=1.\n3. Query history with limit=0, limit=501 and offset=-1.",
        "expected": "Pagination selects the expected conversation/message. Invalid history pagination returns HTTP 422.",
        "actual": "The selected conversation and assistant message matched their offsets; invalid pagination returned HTTP 422.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_history_and_list_pagination_are_bounded"
        ]
    },
    {
        "id": "CH_HIS_04",
        "group": "5. Conversation Listing and Message History APIs",
        "description": "Restore history and continue from a new API client",
        "procedure": "1. Create an owned conversation.\n2. Reconnect with a new API client and the same account JWT.\n3. Restore history and send a follow-up.",
        "expected": "The prior two turns are restored; continuation returns HTTP 201 with the same ID and forwards the restored context.",
        "actual": "A new client restored the pair and continued the same ID with the prior context.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_history_restoration_continues_with_the_same_context_after_new_api_client"
        ]
    },
    {
        "id": "CH_HIS_05",
        "group": "5. Conversation Listing and Message History APIs",
        "description": "Return 404 for a missing conversation",
        "procedure": "1. Use an unknown valid conversation UUID.\n2. Attempt owner-authenticated history retrieval, append and deletion.\n3. Inspect retrieval calls.",
        "expected": "All three operations return HTTP 404 before AI retrieval.",
        "actual": "Read, append and delete returned HTTP 404; retrieval was not invoked.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_missing_conversation_returns_404_on_read_append_and_delete"
        ]
    },
    {
        "id": "CH_OWN_01",
        "group": "6. Ownership Isolation, Authentication and Guest Policy",
        "description": "Block account B from reading account A messages",
        "procedure": "1. Create a conversation as account A.\n2. GET its messages using account B JWT.",
        "expected": "HTTP 403; only a safe error detail is returned.",
        "actual": "Cross-account history access returned HTTP 403 with an error-only response.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_user_b_cannot_read_user_a_messages"
        ]
    },
    {
        "id": "CH_OWN_02",
        "group": "6. Ownership Isolation, Authentication and Guest Policy",
        "description": "Block cross-account conversation continuation before inference",
        "procedure": "1. Create a conversation as A.\n2. POST a new message into its ID using B JWT.\n3. Inspect retrieval and stored row counts.",
        "expected": "HTTP 403 before retrieval; A data remains unchanged.",
        "actual": "The append attempt returned HTTP 403; retrieval was not called and stored data was unchanged.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_user_b_cannot_append_to_user_a_conversation"
        ]
    },
    {
        "id": "CH_OWN_03",
        "group": "6. Ownership Isolation, Authentication and Guest Policy",
        "description": "Block cross-account deletion",
        "procedure": "1. Create a conversation as A.\n2. DELETE its ID using B JWT.\n3. Check existing rows.",
        "expected": "HTTP 403; the conversation, messages and citations remain.",
        "actual": "The delete attempt returned HTTP 403 and the stored conversation data remained.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_user_b_cannot_delete_user_a_conversation"
        ]
    },
    {
        "id": "CH_OWN_04",
        "group": "6. Ownership Isolation, Authentication and Guest Policy",
        "description": "Allow a new guest conversation without Authorization",
        "procedure": "1. POST /api/chat/message without an Authorization header and without a saved ID.\n2. Inspect the stored conversation owner.",
        "expected": "The request succeeds and the conversation account_id is null.",
        "actual": "The guest request succeeded and its stored account_id was null.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_guest_can_create_a_conversation_with_null_account"
        ]
    },
    {
        "id": "CH_OWN_05",
        "group": "6. Ownership Isolation, Authentication and Guest Policy",
        "description": "Do not treat a guest conversation UUID as authentication",
        "procedure": "1. Create a guest conversation.\n2. Attempt a guest append using its saved UUID.\n3. Attempt unauthenticated history retrieval.",
        "expected": "Append returns HTTP 403; unauthenticated history access is rejected by the bearer authentication policy.",
        "actual": "Guest append returned HTTP 403 and unauthenticated history retrieval was rejected.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_guest_cannot_use_a_conversation_uuid_as_authentication"
        ]
    },
    {
        "id": "CH_OWN_06",
        "group": "6. Ownership Isolation, Authentication and Guest Policy",
        "description": "Prevent an account from claiming guest-owned history",
        "procedure": "1. Create a guest conversation.\n2. Authenticate an account and attempt to access that guest conversation.",
        "expected": "HTTP 403; account login does not transfer guest conversation ownership.",
        "actual": "The authenticated account could not claim or access the guest conversation.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_authenticated_account_cannot_claim_guest_history"
        ]
    },
    {
        "id": "CH_OWN_07",
        "group": "6. Ownership Isolation, Authentication and Guest Policy",
        "description": "Require JWT authentication for management APIs",
        "procedure": "1. Without Authorization, request the conversation list, stored history and deletion.\n2. Inspect all responses.",
        "expected": "All management operations reject unauthenticated requests with the configured bearer-scheme status (401/403).",
        "actual": "All unauthenticated management requests were rejected with a permitted bearer-authentication status.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_management_requires_jwt"
        ]
    },
    {
        "id": "CH_OWN_08",
        "group": "6. Ownership Isolation, Authentication and Guest Policy",
        "description": "Reject invalid, expired, reset-purpose and malformed credentials",
        "procedure": "1. Send chat requests with an invalid JWT, expired JWT, password-reset JWT, Basic credentials and an incomplete Bearer header.\n2. Check stored rows.",
        "expected": "Every request returns HTTP 401; none is downgraded to Guest; no rows are stored.",
        "actual": "All invalid authorization variants returned HTTP 401 with no stored conversation data.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_invalid_expired_wrong_purpose_and_malformed_authorization_are_not_guests"
        ]
    },
    {
        "id": "CH_OWN_09",
        "group": "6. Ownership Isolation, Authentication and Guest Policy",
        "description": "Reject deleted or inactive account identities",
        "procedure": "1. Use validly signed tokens referencing an inactive account and an unknown account.\n2. POST /api/chat/message.",
        "expected": "HTTP 401 for both identities; no conversation data stored.",
        "actual": "Both account states were rejected with HTTP 401 and no conversation data was stored.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_deleted_and_inactive_accounts_are_rejected"
        ]
    },
    {
        "id": "CH_OWN_10",
        "group": "6. Ownership Isolation, Authentication and Guest Policy",
        "description": "Keep bounded model history isolated between accounts",
        "procedure": "1. Create A and B conversations with distinct private content.\n2. Continue A repeatedly.\n3. Inspect model history and attempt access to B through A.",
        "expected": "Only recent A turns reach its model context; B content never appears; unauthorized inference is blocked.",
        "actual": "The six-message context contained no B content or outdated first turn; access to B was rejected before generation.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_previous_messages_are_bounded_and_isolated_between_accounts"
        ]
    },
    {
        "id": "CH_RES_01",
        "group": "7. Inference Timeouts, Failures and Development Fallback",
        "description": "Roll back a timed-out new conversation",
        "procedure": "1. In the isolated non-fallback test, simulate slow retrieval.\n2. Trigger the configured overall chat deadline using a shortened test clock.\n3. Check response and stored rows.",
        "expected": "The configured deadline is 30 seconds; timeout returns HTTP 504 with a safe detail and no partial conversation data.",
        "actual": "The 30-second policy and HTTP 504 were verified; no conversation, message or citation remained.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_timeout_returns_504_and_does_not_leave_partial_conversation"
        ]
    },
    {
        "id": "CH_RES_02",
        "group": "7. Inference Timeouts, Failures and Development Fallback",
        "description": "Preserve existing history when continuation times out",
        "procedure": "1. Create an owned conversation.\n2. Simulate a timeout while appending.\n3. Compare history, citations and updated_at before/after.",
        "expected": "HTTP 504; existing turns, citations and updated_at remain unchanged.",
        "actual": "The timed-out continuation returned HTTP 504 and preserved previous rows and updated_at.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_timeout_on_existing_conversation_preserves_previous_turns"
        ]
    },
    {
        "id": "CH_RES_03",
        "group": "7. Inference Timeouts, Failures and Development Fallback",
        "description": "Prevent late inference output from writing after timeout",
        "procedure": "1. Make generation continue in a worker after the request deadline.\n2. Wait for that controlled worker to finish.\n3. Check the database.",
        "expected": "HTTP 504; late inference completion cannot persist a conversation or message.",
        "actual": "The late worker finished after HTTP 504; no conversation, message or citation was written.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_inference_timeout_cannot_write_a_late_answer"
        ]
    },
    {
        "id": "CH_RES_04",
        "group": "7. Inference Timeouts, Failures and Development Fallback",
        "description": "Handle primary model failure safely",
        "procedure": "1. In non-fallback mode, make primary inference raise an internal error.\n2. Send a persisted chat request.\n3. Inspect the response and database.",
        "expected": "HTTP 500 with a safe detail; internal provider information is absent; no partial data stored.",
        "actual": "The model failure returned HTTP 500 without private details and rolled back all new rows.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_model_failure_returns_safe_500_and_rolls_back"
        ]
    },
    {
        "id": "CH_RES_05",
        "group": "7. Inference Timeouts, Failures and Development Fallback",
        "description": "Handle retrieval failure safely",
        "procedure": "1. In non-fallback mode, make evidence retrieval fail.\n2. Send a persisted request.\n3. Inspect its response and database.",
        "expected": "HTTP 500 with a safe detail and no partial persisted data.",
        "actual": "The retrieval failure returned HTTP 500, omitted private details and left no new rows.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_retrieval_failure_returns_safe_500_and_rolls_back"
        ]
    },
    {
        "id": "CH_RES_06",
        "group": "7. Inference Timeouts, Failures and Development Fallback",
        "description": "Roll back a failed commit after messages and citations were flushed",
        "procedure": "1. Make the isolated database commit raise an error after flush.\n2. Send a valid chat question.\n3. Inspect the response and database.",
        "expected": "HTTP 500; all flushed conversation/message/citation rows are rolled back; private database details are not exposed.",
        "actual": "The failed commit returned HTTP 500, omitted private details and left no partially stored rows.",
        "tests": [
            "apps.backend.tests.test_chat_conversation.ChatConversationTests::test_commit_failure_rolls_back_messages_and_citations_already_flushed"
        ]
    },
    {
        "id": "CH_RES_07",
        "group": "7. Inference Timeouts, Failures and Development Fallback",
        "description": "Enforce the compatibility interface inference deadline",
        "procedure": "1. Simulate slow compatibility-route retrieval.\n2. Trigger its shortened test deadline.\n3. Inspect the error detail.",
        "expected": "HTTP 504; the detail refers to the same overall 30-second policy.",
        "actual": "The compatibility request returned HTTP 504 with the 30-second policy detail.",
        "tests": [
            "apps.backend.tests.test_chat_api.ChatApiTests::test_legacy_interface_has_the_same_30_second_inference_deadline"
        ]
    },
    {
        "id": "CH_RES_08",
        "group": "7. Inference Timeouts, Failures and Development Fallback",
        "description": "Report compatibility retrieval unavailability without fabricated answers",
        "procedure": "1. Make compatibility retrieval fail its source integrity check.\n2. POST /api/chat.\n3. Inspect resolution and sources.",
        "expected": "The response declares resolution_status=unavailable and returns no sources.",
        "actual": "The compatibility response declared unavailable retrieval with an empty source list.",
        "tests": [
            "apps.backend.tests.test_chat_api.ChatApiTests::test_storage_failure_is_not_reported_as_an_answer"
        ]
    },
    {
        "id": "CH_RES_09",
        "group": "7. Inference Timeouts, Failures and Development Fallback",
        "description": "Restrict development fallback to local mode",
        "procedure": "1. Set the environment to production, staging, test and an unknown value in separate isolated runs.\n2. Attempt the development fallback.",
        "expected": "Fallback is disabled and direct fallback invocation is rejected in each non-local environment.",
        "actual": "Every non-local environment disabled fallback and rejected direct invocation.",
        "tests": [
            "apps.backend.tests.test_chat_fallback.DevelopmentAnswerTests::test_only_local_environment_enables_fallback"
        ]
    },
    {
        "id": "CH_RES_10",
        "group": "7. Inference Timeouts, Failures and Development Fallback",
        "description": "Persist a labelled local guest fallback when retrieval is unavailable",
        "procedure": "1. Enable local development mode without an AI key.\n2. Simulate retrieval connection failure.\n3. Send a guest question and inspect persisted records.",
        "expected": "HTTP 201; answer_type=development_mock; the new pair and demo citation are stored. Demo citation URL is null.",
        "actual": "The labelled mock persisted one conversation, two turns and a demo citation with a null URL; primary generation was not called.",
        "tests": [
            "apps.backend.tests.test_chat_fallback.LocalConversationFallbackTests::test_unavailable_retrieval_persists_guest_pair_and_demo_citation"
        ]
    },
    {
        "id": "CH_RES_11",
        "group": "7. Inference Timeouts, Failures and Development Fallback",
        "description": "Keep retrieved sources when the local model is missing",
        "procedure": "1. Retrieve valid evidence in local mode.\n2. Simulate a missing GGUF/model file.\n3. Inspect fallback history, citations and access by another account.",
        "expected": "A labelled mock preserves the real source evidence/citation and owned history; another account is denied.",
        "actual": "The fallback retained source evidence and URL, persisted the pair and rejected account B history access with HTTP 403.",
        "tests": [
            "apps.backend.tests.test_chat_fallback.LocalConversationFallbackTests::test_missing_llm_or_gguf_keeps_retrieved_sources_and_account_history"
        ]
    },
    {
        "id": "CH_RES_12",
        "group": "7. Inference Timeouts, Failures and Development Fallback",
        "description": "Persist successful Gemini fallback with original citations",
        "procedure": "1. Simulate an unavailable primary model in local mode.\n2. Supply a controlled successful Gemini response with retrieved evidence.\n3. Inspect persisted answer and citations.",
        "expected": "The Gemini answer is stored with resolution_status=gemini_fallback and the original source citation.",
        "actual": "The Gemini answer, fallback status and original citation were stored as one conversation and message pair.",
        "tests": [
            "apps.backend.tests.test_chat_fallback.LocalConversationFallbackTests::test_gemini_success_is_saved_with_original_citations"
        ]
    },
    {
        "id": "CH_RES_13",
        "group": "7. Inference Timeouts, Failures and Development Fallback",
        "description": "Recover from Gemini errors and timeouts without leaking secrets",
        "procedure": "1. Simulate connection, quota, malformed-response and parsing failures at the Gemini boundary.\n2. Simulate its bounded timeout.\n3. Inspect the returned answer.",
        "expected": "A labelled development mock is returned; no API key, internal configuration or upstream private body is exposed.",
        "actual": "Every controlled Gemini failure and timeout produced a labelled mock without secret/private error details.",
        "tests": [
            "apps.backend.tests.test_chat_fallback.DevelopmentAnswerTests::test_gemini_errors_or_malformed_responses_fall_back_without_leaking_details",
            "apps.backend.tests.test_chat_fallback.DevelopmentAnswerTests::test_gemini_timeout_is_bounded_and_returns_mock"
        ]
    },
    {
        "id": "CH_RES_14",
        "group": "7. Inference Timeouts, Failures and Development Fallback",
        "description": "Keep the healthy primary path independent of fallback",
        "procedure": "1. Enable local mode with a configured test key.\n2. Make primary retrieval/inference succeed.\n3. Check fallback calls.",
        "expected": "The result is answered by the primary path; Gemini/mock fallback is not invoked.",
        "actual": "The result was classified as answered and the fallback was not called.",
        "tests": [
            "apps.backend.tests.test_chat_fallback.LocalConversationFallbackTests::test_healthy_primary_does_not_call_gemini_or_mock"
        ]
    },
    {
        "id": "CH_INT_01",
        "group": "8. Chat API Integration and History Restoration",
        "description": "Exercise JWT, retrieval assembly, provider HTTP parsing and restoration together",
        "procedure": "1. Create A and B conversations using the real router, JWT validation, SQLAlchemy and LLM HTTP adapter with controlled retrieval/provider responses.\n2. Reconnect as A, restore history and continue.\n3. Inspect the provider request.",
        "expected": "History and citations restore correctly; continuation returns HTTP 201. Only A history enters the provider request; B cannot read A history.",
        "actual": "The full assembled flow restored citations and context, returned HTTP 201 and excluded B private content from A provider input.",
        "tests": [
            "apps.backend.tests.test_chat_integration.ChatPipelineIntegrationTests::test_account_history_restoration_and_continuation_cross_the_provider_boundary"
        ]
    },
    {
        "id": "CH_INT_02",
        "group": "8. Chat API Integration and History Restoration",
        "description": "Exercise embedded compatibility chat without creating persisted history",
        "procedure": "1. POST /api/chat with heritage, scene and prior context using the real provider adapter with a controlled HTTP response.\n2. Inspect answer, sources, context and stored row counts.",
        "expected": "HTTP 200 with answer/sources and context-aware model input; the stateless compatibility request does not create conversation rows.",
        "actual": "The compatibility response and context-aware provider request were verified; conversation, message and citation counts remained zero.",
        "tests": [
            "apps.backend.tests.test_chat_integration.ChatPipelineIntegrationTests::test_embedded_compatibility_chat_has_context_citations_and_no_persistence_requirement"
        ]
    },
    {
        "id": "CH_INT_03",
        "group": "8. Chat API Integration and History Restoration",
        "description": "Reject empty, truncated or filtered provider output",
        "procedure": "1. Make the real HTTP adapter parse empty content, length-truncated content and filtered output.\n2. Call persisted and compatibility interfaces.\n3. Inspect stored rows.",
        "expected": "Persisted requests return HTTP 500; compatibility requests return HTTP 503; unusable answers are not committed.",
        "actual": "All controlled incomplete outputs were rejected with the expected statuses and no new stored rows.",
        "tests": [
            "apps.backend.tests.test_chat_integration.ChatPipelineIntegrationTests::test_empty_or_truncated_inference_is_not_committed_or_returned_as_an_answer"
        ]
    },
    {
        "id": "CH_INT_04",
        "group": "8. Chat API Integration and History Restoration",
        "description": "Handle provider connection failure consistently across both interfaces",
        "procedure": "1. Make the real provider HTTP call raise a connection error containing private diagnostic text.\n2. Call both chat interfaces in non-fallback mode.\n3. Check error responses and stored rows.",
        "expected": "Persisted HTTP 500 and compatibility HTTP 503; private diagnostic text is absent; no partial persistence.",
        "actual": "Both interfaces returned their safe expected failure statuses with no private text or stored partial data.",
        "tests": [
            "apps.backend.tests.test_chat_integration.ChatPipelineIntegrationTests::test_provider_connection_failure_is_safe_and_atomic_for_both_interfaces"
        ]
    },
    {
        "id": "CH_INT_05",
        "group": "8. Chat API Integration and History Restoration",
        "description": "Load environment configuration before authentication imports",
        "procedure": "1. In an isolated subprocess, supply a controlled dotenv configuration before importing the app.\n2. Inspect the JWT configuration seen by the security module.",
        "expected": "Authentication reads the configured signing secret after dotenv loading.",
        "actual": "The isolated subprocess confirmed that environment configuration was loaded before authentication initialization.",
        "tests": [
            "apps.backend.tests.test_chat_integration.ChatStartupTests::test_env_configuration_is_loaded_before_authentication_modules"
        ]
    },
    {
        "id": "CH_INT_06",
        "group": "8. Chat API Integration and History Restoration",
        "description": "Prevent history from replacing the model system policy",
        "procedure": "1. Pass system, developer and tool roles as prior history directly to prompt assembly in separate isolated runs.\n2. Inspect validation.",
        "expected": "Prompt assembly rejects each non-user/assistant history role.",
        "actual": "All disallowed history roles raised validation errors instead of overriding the system prompt.",
        "tests": [
            "apps.backend.tests.test_prompt.OverviewPromptTests::test_history_cannot_override_the_system_prompt"
        ]
    }
]


def read_evidence(path: Path | None):
    if path is None:
        return {}, "", "No automated execution evidence supplied.", 0
    raw = Path(path).read_bytes()
    root = ElementTree.fromstring(raw)
    outcomes = defaultdict(list)
    executed_nodes = 0
    for node in root.iter("testcase"):
        classname = node.get("classname", "")
        if classname not in TEST_CLASSES:
            continue
        executed_nodes += 1
        method = node.get("name", "").split("[", 1)[0].strip()
        reference = f"{classname}::{method}"
        if node.find("failure") is not None or node.find("error") is not None:
            result = "Fail"
        elif node.find("skipped") is not None:
            result = "Not Run"
        else:
            result = "Pass"
        outcomes[reference].append(result)
    if not outcomes:
        raise ValueError("JUnit does not contain any of the existing Chat Backend test classes.")
    states = {
        reference: "Fail" if "Fail" in values else "Not Run" if "Not Run" in values else "Pass"
        for reference, values in outcomes.items()
    }
    suite = next(root.iter("testsuite"))
    executed_at = datetime.fromisoformat(suite.attrib["timestamp"])
    if executed_at.tzinfo is None:
        executed_at = executed_at.replace(tzinfo=timezone.utc)
    date = executed_at.astimezone(timezone(timedelta(hours=7))).strftime("%d/%m/%Y")
    evidence = f"JUnit file: {Path(path).name}; SHA256: {hashlib.sha256(raw).hexdigest()}."
    return states, date, evidence, executed_nodes


def case_result(case, states):
    results = [states.get(reference, "Not Run") for reference in case["tests"]]
    if "Fail" in results:
        return "Fail"
    if not results or "Not Run" in results:
        return "Not Run"
    return "Pass"


def build_report(output=OUTPUT, junit=None):
    if len(CASES) != len({case["id"] for case in CASES}):
        raise ValueError("FUT case IDs must be unique.")
    if any(not case["tests"] for case in CASES):
        raise ValueError("Each FUT case must map to executed test evidence.")
    states, date, evidence, executed_count = read_evidence(junit)
    counts = Counter(case_result(case, states) for case in CASES)
    book = Workbook()
    sheet = book.active
    sheet.title = "Chat_Conversation_Management"
    book.properties.creator = TESTER
    book.properties.lastModifiedBy = TESTER
    book.properties.title = "FUT - Chat Backend and Account-Based Conversation Management"
    book.properties.subject = MODULE
    book.properties.description = f"Backend source {SOURCE_REVISION}. {evidence} {MANUAL_QA}"
    for col, width in enumerate(WIDTHS, 1):
        sheet.column_dimensions[chr(64 + col)].width = width

    metadata = (
        ("Module Code", MODULE),
        ("Tester", TESTER),
        (
            "Pre-condition",
            "Manual API checks: FastAPI http://127.0.0.1:8000; migrated heritagegraph database; "
            "active accounts A/B; published heritage data; dedicated test conversations.\n"
            "Automated checks: isolated SQLite, real ORM/JWT handling, controlled retrieval and provider "
            "HTTP responses. Fault tests disable local fallback unless fallback itself is under test.",
        ),
    )
    for row, (label, value) in enumerate(metadata, 1):
        cell = sheet.cell(row, 1, label)
        cell.font = font(bold=True)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        merged(sheet, row, 2, 7, value, centered=row < 3)
        sheet.row_dimensions[row].height = 26 if row < 3 else 66
    sheet["A1"].comment = Comment(
        f"Task 2 source revision: {SOURCE_REVISION}. "
        "Scope: both chat interfaces, AI outcomes, heritage/scene/bounded context, persistence, "
        "listing/history, ownership/guests, resilience and integration.\n"
        f"{MANUAL_QA}\nAutomated fixture coverage does not assess live model/corpus answer quality.",
        TESTER,
    )

    for address, label in (("A4", "Pass"), ("B4", "Fail")):
        cell = sheet[address]
        cell.value = label
        cell.font = font(bold=True, size=12)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    merged(sheet, 4, 3, 7, "Number of test case", bold=True, centered=True)
    for address, value in (("A5", counts["Pass"]), ("B5", counts["Fail"])):
        cell = sheet[address]
        cell.value = value
        cell.font = font()
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    merged(sheet, 5, 3, 7, len(CASES), centered=True)
    sheet["C5"].comment = Comment(
        f"{counts['Pass']} Pass + {counts['Fail']} Fail + {counts['Not Run']} Not Run = {len(CASES)}. "
        f"Matching JUnit test records: {executed_count}. "
        "FUT cases may combine several methods or share scenario evidence; counts are not additive. "
        "Regenerate after another execution.",
        TESTER,
    )
    sheet.row_dimensions[4].height = 25
    sheet.row_dimensions[5].height = 25
    sheet.row_dimensions[6].height = 12

    for col, title in enumerate(HEADERS, 1):
        cell = sheet.cell(7, col, title)
        cell.font = font(bold=True, size=12)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[7].height = 36
    sheet["F7"].comment = Comment(
        "Pass requires every mapped method/subtest to pass. Expected negative HTTP responses "
        "also count as successful tests. A missing or skipped dependency makes the case Not Run; "
        "any failed dependency makes it Fail. " + evidence,
        TESTER,
    )
    dropdown = DataValidation(type="list", formula1='"Pass,Fail,Not Run"', allow_blank=False)
    dropdown.errorTitle = "Invalid result"
    dropdown.error = "Choose Pass, Fail or Not Run."
    dropdown.showErrorMessage = True
    dropdown.errorStyle = "stop"
    sheet.add_data_validation(dropdown)

    row, group = 8, None
    for case in CASES:
        if case["group"] != group:
            group = case["group"]
            merged(sheet, row, 1, 7, group, bold=True, color=ORANGE, centered=True)
            sheet.row_dimensions[row].height = 28
            row += 1
        result = case_result(case, states)
        actual = case["actual"] if result == "Pass" else (
            "At least one mapped test failed. Review its JUnit failure details."
            if result == "Fail" else "Required test evidence is missing or skipped. Execute before sign-off."
        )
        values = [
            case["id"], case["description"], case["procedure"], case["expected"],
            actual, result, date if result != "Not Run" else "",
        ]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row, col)
            cell.font = font(bold=col in (1, 6))
            cell.border = BORDER
            cell.alignment = Alignment(
                horizontal="center" if col in (1, 6, 7) else "left",
                vertical="center" if col in (1, 6, 7) else "top", wrap_text=True,
            )
            cell.value = datetime.strptime(value, "%d/%m/%Y") if col == 7 and value else value or None
            if col == 7:
                cell.number_format = "dd/mm/yyyy"
            if col == 6:
                fill, text = RESULT_COLORS[result]
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = font(bold=True, color=text)
                dropdown.add(cell)
                for state, (bg, fg) in RESULT_COLORS.items():
                    sheet.conditional_formatting.add(
                        cell.coordinate, CellIsRule(
                            operator="equal", formula=[f'"{state}"'],
                            fill=PatternFill("solid", fgColor=bg), font=font(bold=True, color=fg),
                        ),
                    )
        sheet.cell(row, 5).comment = Comment(
            "Mapped automated evidence:\n" + "\n".join(case["tests"]) +
            f"\nExecution date (Asia/Saigon): {date or 'Not Run'}. {evidence}\n"
            "References correspond to the existing test modules under apps/backend/tests/. "
            "Controlled retrieval/provider responses exercise deterministic success and error paths.",
            TESTER,
        )
        sheet.row_dimensions[row].height = estimated_height(values, WIDTHS)
        row += 1

    sheet.auto_filter.ref = f"A7:G{row - 1}"
    notes = (
        f"Summary: {len(CASES)} FUT cases | {counts['Pass']} Pass | {counts['Fail']} Fail | "
        f"{counts['Not Run']} Not Run. Automated evidence contains {executed_count} matching test records.",
        MANUAL_QA,
        "Automated fixtures exercise the existing router, persistence, JWT validation, retrieval assembly "
        "and provider adapters with controlled external responses. This is not a live AI answer-quality evaluation.",
        "The local labelled mock/Gemini fallback is tested separately from non-local 500/504 failure paths. "
        "Generating this workbook does not execute tests or change the database.",
    )
    row += 1
    for note in notes:
        merged(sheet, row, 1, 7, note)
        sheet.row_dimensions[row].height = 36
        row += 1
    sheet.freeze_panes = "B8"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 75
    sheet.print_area = f"A1:G{row - 1}"
    sheet.print_title_rows = "7:7"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4)
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.left.text = TESTER
    sheet.oddFooter.right.text = "Page &P / &N"
    book.save(output)
    book.close()
    return counts


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--junit", type=Path, help="pytest JUnit XML from the existing Chat test modules")
    parser.add_argument("--output", type=Path, default=OUTPUT, help="Destination workbook")
    args = parser.parse_args()
    counts = build_report(args.output, args.junit)
    print(f"Created: {args.output}")
    print(f"Pass: {counts['Pass']} | Fail: {counts['Fail']} | Not Run: {counts['Not Run']} | Total: {len(CASES)}")
    return 1 if counts["Fail"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
