"""Generate the tester-supplied Account Management QA report.

This script formats the supplied results; it does not run application tests.
Run with: python generate_account_test.py
Requires: openpyxl
"""
from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path
from textwrap import wrap

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

TESTER = "Cao Văn Khoa"
OUTPUT_PATH = Path(__file__).resolve().with_name("Testcase_AccountManagement.xlsx")
COLUMNS = (
    "ID", "Test case Description", "Test case Procedure", "Expected Result",
    "Actual Result", "Result", "Test Date",
)
COLUMN_WIDTHS = (12, 35, 45, 35, 35, 10, 13)
SECTIONS = [
    {
        "title": "REGISTER",
        "cases": [
            [
                "AM_REG_01",
                "Register with valid email, username and strong password",
                "1. Navigate to Register tab\n2. Enter valid email (khoa_qa@example.com), username (khoa_qa) and strong password (Password@2026)\n3. Click \"Đăng ký\"",
                "Account created successfully; HTTP 201 Created; User redirected to Login page",
                "Account created; HTTP 201; Redirected to Login",
                "Pass",
                "28/10/2026"
            ],
            [
                "AM_REG_02",
                "Register with an already existing email",
                "1. Enter an email already registered\n2. Complete other fields\n3. Click \"Đăng ký\"",
                "Error message: \"Email đã tồn tại trong hệ thống.\"; HTTP 400",
                "Error message: \"Email đã tồn tại trong hệ thống.\"",
                "Pass",
                "28/10/2026"
            ],
            [
                "AM_REG_03",
                "Register with an already existing username",
                "1. Enter a username already taken in DB\n2. Complete other fields\n3. Click \"Đăng ký\"",
                "Error message: \"Tên đăng nhập đã được sử dụng.\"; HTTP 400",
                "Error message: \"Tên đăng nhập đã được sử dụng.\"",
                "Pass",
                "28/10/2026"
            ],
            [
                "AM_REG_04",
                "Register with invalid email format",
                "Enter \"user@com\" or \"user123\" and click \"Đăng ký\"",
                "Validation error: \"Định dạng email không hợp lệ.\"; HTTP 422",
                "Validation error: \"Định dạng email không hợp lệ.\"",
                "Pass",
                "28/10/2026"
            ],
            [
                "AM_REG_05",
                "Register with weak password (< 8 chars)",
                "Enter password \"12345\" and click \"Đăng ký\"",
                "Validation error: \"Mật khẩu phải chứa ít nhất 8 ký tự.\"; HTTP 422",
                "Validation error displayed",
                "Pass",
                "28/10/2026"
            ],
            [
                "AM_REG_06",
                "Register with empty required fields",
                "Leave Email, Username or Password blank and click \"Đăng ký\"",
                "Form blocks submission; Alert box highlights required fields",
                "Form blocks submission; required fields highlighted",
                "Pass",
                "28/10/2026"
            ]
        ]
    },
    {
        "title": "EMAIL VERIFICATION",
        "cases": [
            [
                "AM_VER_01",
                "Verify account via valid email link",
                "Click the activation link sent to the registered email",
                "Redirect to Login page with \"Verification Successful\" message",
                "Email verification feature is not yet available in Sprint 1 (Backlog)",
                "Fail",
                "28/10/2026"
            ],
            [
                "AM_VER_02",
                "Use an expired verification link",
                "Click a verification link that has passed its expiry time",
                "Error message: \"Link expired. Please request a new one.\"",
                "Email verification feature is not yet available in Sprint 1 (Backlog)",
                "Fail",
                "28/10/2026"
            ]
        ]
    },
    {
        "title": "LOGIN",
        "cases": [
            [
                "AM_LOG_01",
                "Login with registered email and correct password",
                "1. Enter registered email (khoa_qa@example.com) and correct password\n2. Click \"Đăng nhập\"",
                "Successful login; JWT token issued; Redirected to Homepage",
                "Successful login; Session established",
                "Pass",
                "28/10/2026"
            ],
            [
                "AM_LOG_02",
                "Login with username instead of email",
                "1. Enter username (khoa_qa) and correct password\n2. Click \"Đăng nhập\"",
                "Successful login; HTTP 200 OK; JWT returned",
                "Successful login; Token issued",
                "Pass",
                "28/10/2026"
            ],
            [
                "AM_LOG_03",
                "Login with incorrect password",
                "1. Enter registered email and wrong password\n2. Click \"Đăng nhập\"",
                "Error message: \"Email hoặc mật khẩu không chính xác.\"; HTTP 401",
                "Error message displayed",
                "Pass",
                "28/10/2026"
            ],
            [
                "AM_LOG_04",
                "Login with non-existent email",
                "1. Enter unregistered email notfound@example.com\n2. Click \"Đăng nhập\"",
                "Error message: \"Tài khoản không tồn tại.\"; HTTP 401",
                "Error message displayed",
                "Pass",
                "28/10/2026"
            ],
            [
                "AM_LOG_05",
                "Login with empty credentials",
                "Leave Email and Password blank and click \"Đăng nhập\"",
                "Validation error: \"Vui lòng điền đầy đủ email và mật khẩu.\"",
                "Validation error displayed on UI",
                "Pass",
                "28/10/2026"
            ]
        ]
    },
    {
        "title": "GOOGLE OAUTH 2.0",
        "cases": [
            [
                "AM_OAU_01",
                "First-time login with Google OAuth",
                "1. Click \"Đăng nhập với Google\"\n2. Authenticate with Gmail account",
                "Account auto-created; Google avatar (picture) pulled; JWT session returned",
                "Account created; Avatar fetched from Google",
                "Pass",
                "28/10/2026"
            ],
            [
                "AM_OAU_02",
                "Subsequent login with existing Google account",
                "1. Click \"Đăng nhập với Google\" with previously linked Gmail",
                "Direct login success; Profile avatar synced; Redirect to Homepage",
                "Direct login success; Token issued",
                "Pass",
                "28/10/2026"
            ],
            [
                "AM_OAU_03",
                "Cancel Google OAuth popup",
                "Click Google Login then close the popup dialog",
                "Auth flow aborted; User stays on Login page with no error",
                "Auth flow aborted; Remains on Login page",
                "Pass",
                "28/10/2026"
            ]
        ]
    },
    {
        "title": "AUTHENTICATION & LOGOUT",
        "cases": [
            [
                "AM_AUT_01",
                "Access protected API with valid JWT",
                "Send GET /api/profile/me with \"Authorization: Bearer <valid_token>\"",
                "HTTP 200 OK; User profile data returned",
                "HTTP 200 OK; Returns profile",
                "Pass",
                "28/10/2026"
            ],
            [
                "AM_AUT_02",
                "Access protected API without JWT",
                "Send GET /api/profile/me without Authorization header",
                "HTTP 401 Unauthorized; Message: \"Not authenticated\"",
                "HTTP 401 Unauthorized",
                "Pass",
                "28/10/2026"
            ],
            [
                "AM_AUT_03",
                "User Logout functionality",
                "1. Click \"Đăng xuất\" on User Menu\n2. Try to access /profile",
                "JWT token cleared; Auth context reset; User redirected to /login",
                "Token cleared; Redirect to /login",
                "Pass",
                "28/10/2026"
            ]
        ]
    },
    {
        "title": "PASSWORD MANAGEMENT",
        "cases": [
            [
                "AM_PWD_01",
                "Change password with correct old password",
                "1. In Profile, enter correct old password and new strong password\n2. Click \"Đổi mật khẩu\"",
                "Password updated successfully; HTTP 200 OK",
                "Password updated; New password works",
                "Pass",
                "28/10/2026"
            ],
            [
                "AM_PWD_02",
                "Change password with incorrect old password",
                "1. Enter wrong current password\n2. Click \"Đổi mật khẩu\"",
                "Error message: \"Mật khẩu hiện tại không chính xác.\"; HTTP 400",
                "Error message displayed",
                "Pass",
                "28/10/2026"
            ]
        ]
    }
]

THIN_BLACK = Side(style="thin", color="FF000000")
CELL_BORDER = Border(left=THIN_BLACK, right=THIN_BLACK, top=THIN_BLACK, bottom=THIN_BLACK)
HEADER_FILL = PatternFill("solid", fgColor="FFD9E1F2")
TABLE_FILL = PatternFill("solid", fgColor="FF1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="FFED7D31")
RESULT_STYLES = {
    "Pass": ("FFC6EFCE", "FF006100"),
    "Fail": ("FFFFC7CE", "FF9C0006"),
}


def estimated_row_height(values, widths) -> float:
    """Estimate wrapped lines because openpyxl cannot calculate Excel AutoFit."""
    line_count = 1
    for value, width in zip(values, widths):
        usable_width = max(1, int(width - 2))
        lines = str(value).split("\n")
        wrapped_lines = sum(
            max(1, len(wrap(line, width=usable_width, break_long_words=True)))
            for line in lines
        )
        line_count = max(line_count, wrapped_lines)
    return max(40, 16 * line_count + 12)


def merged_row(sheet, row, first_col, last_col, value, fill, font, alignment):
    sheet.merge_cells(
        start_row=row, start_column=first_col, end_row=row, end_column=last_col,
    )
    sheet.cell(row, first_col, value)
    for col in range(first_col, last_col + 1):
        cell = sheet.cell(row, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment
        cell.border = CELL_BORDER


def build_report(output_path: Path = OUTPUT_PATH) -> Path:
    cases = [case for section in SECTIONS for case in section["cases"]]
    counts = Counter(case[5] for case in cases)
    if counts != Counter({"Pass": 19, "Fail": 2}) or len(cases) != 21:
        raise ValueError("The supplied report must contain 19 Pass and 2 Fail cases.")
    if len({case[0] for case in cases}) != len(cases):
        raise ValueError("Test case IDs must be unique.")
    if any(len(case) != len(COLUMNS) for case in cases):
        raise ValueError("Every test case must have seven fields.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Account_Management"
    workbook.properties.creator = TESTER
    workbook.properties.lastModifiedBy = TESTER
    workbook.properties.title = "Task 1 - Account Management Test Report"
    workbook.properties.subject = "Account Management (Auth & Profile)"
    workbook.properties.description = (
        "QA results supplied by Cao Văn Khoa for 28/10/2026: "
        "19 Pass, 2 Fail, 21 total. The generator does not execute tests."
    )

    regular = Font(name="Calibri", size=11, color="FF000000")
    bold = Font(name="Calibri", size=11, bold=True, color="FF000000")
    white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width

    header_values = (
        ("Module Code", "Account Management (Auth & Profile)"),
        ("Tester", TESTER),
        (
            "Pre-condition",
            "Internet active, Backend API (FastAPI Port 8000), Frontend (Port 3000), "
            "Database heritagegraph loaded.",
        ),
        ("Pass / Fail / Total", f"{counts['Pass']} / {counts['Fail']} / {len(cases)}"),
    )
    for row, (label, value) in enumerate(header_values, start=1):
        label_cell = sheet.cell(row, 1, label)
        label_cell.font = bold
        label_cell.fill = HEADER_FILL
        label_cell.border = CELL_BORDER
        label_cell.alignment = left
        merged_row(sheet, row, 2, 7, value, HEADER_FILL, bold, left)
        sheet.row_dimensions[row].height = estimated_row_height(
            (label, value), (COLUMN_WIDTHS[0], sum(COLUMN_WIDTHS[1:])),
        )

    sheet.row_dimensions[5].height = 12
    for col, value in enumerate(COLUMNS, start=1):
        cell = sheet.cell(6, col, value)
        cell.font = white_bold
        cell.fill = TABLE_FILL
        cell.alignment = centered
        cell.border = CELL_BORDER
    sheet.row_dimensions[6].height = 40
    sheet["F6"].comment = Comment(
        "Results supplied by Cao Văn Khoa. This report generator does not execute tests.",
        TESTER,
    )

    row = 7
    for section in SECTIONS:
        merged_row(sheet, row, 1, 7, section["title"], SECTION_FILL, white_bold, centered)
        sheet.row_dimensions[row].height = 40
        row += 1
        for case in section["cases"]:
            for col, value in enumerate(case, start=1):
                cell = sheet.cell(row, col)
                cell.border = CELL_BORDER
                cell.font = bold if col in (1, 6) else regular
                cell.alignment = centered if col in (1, 6, 7) else left
                if col == 7:
                    cell.value = datetime.strptime(value, "%d/%m/%Y")
                    cell.number_format = "dd/mm/yyyy"
                else:
                    cell.value = value
                if col == 6:
                    background, foreground = RESULT_STYLES[value]
                    cell.fill = PatternFill("solid", fgColor=background)
                    cell.font = Font(name="Calibri", size=11, bold=True, color=foreground)
            sheet.row_dimensions[row].height = estimated_row_height(case, COLUMN_WIDTHS)
            row += 1

    sheet.freeze_panes = "B7"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 80
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4)
    sheet.print_area = f"A1:G{row - 1}"
    sheet.print_title_rows = "6:6"
    sheet.oddFooter.left.text = TESTER
    sheet.oddFooter.right.text = "Page &P / &N"
    output_path = Path(output_path)
    workbook.save(output_path)
    workbook.close()
    return output_path


if __name__ == "__main__":
    print(f"Created: {build_report()}")
    print(f"Tester: {TESTER} | Pass / Fail / Total: 19 / 2 / 21")

