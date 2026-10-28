"""Generate the English Backend Login/Register FUT from pytest JUnit evidence.

Run the existing auth tests with --junitxml=auth-fut.xml, then:
    python generate_auth_fut.py --junit auth-fut.xml

Without JUnit evidence, cases remain Not Run. This script does not execute APIs.
"""
from __future__ import annotations

import argparse
import hashlib
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from textwrap import wrap
from xml.etree import ElementTree

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

TESTER = "Cao Văn Khoa"
OUTPUT = Path(__file__).resolve().with_name("FUT_BE_Login_Register.xlsx")
SOURCE_REVISION = "294e3bf"
HEADERS = (
    "ID", "Test case Description", "Test case Procedure", "Expected Result",
    "Actual Result", "Result", "Test Date",
)
WIDTHS = (16, 40, 60, 48, 44, 12, 14)
CASES = [
    {
        "id": "AM_REG_01",
        "group": "Register",
        "description": "Register a valid account and persist its credentials securely",
        "procedure": "1. Use a new email and username.\n2. POST /api/auth/register with email=\"fut_khoa@example.com\", username=\"fut_khoa\", password=\"Password@2026\".\n3. Inspect the response and account row.",
        "expected": "HTTP 201; string account_id, email, username and message=\"Registration successful\". Database role_id=1, is_active=true, timestamps set; Argon2 hash verifies the password. No password/hash in the response.",
        "actual": "HTTP 201; account fields, USER role, active status, timestamps and Argon2 password verification passed. No password/hash was returned.",
        "test": "test_register_creates_hashed_user_account_with_timestamps"
    },
    {
        "id": "AM_REG_02",
        "group": "Register",
        "description": "Reject an already registered email",
        "procedure": "1. Register an account.\n2. Repeat POST /api/auth/register with its email and a different username.\n3. Compare account counts.",
        "expected": "HTTP 400; detail=\"This email is already in use.\"; account count unchanged.",
        "actual": "HTTP 400; email conflict returned; no additional account was stored.",
        "test": "test_register_duplicate_email_returns_400_without_another_account"
    },
    {
        "id": "AM_REG_03",
        "group": "Register",
        "description": "Reject an already registered username",
        "procedure": "1. Register an account.\n2. POST /api/auth/register using the same username and a new email.\n3. Compare account counts.",
        "expected": "HTTP 400; detail=\"This username is already taken.\"; account count unchanged.",
        "actual": "HTTP 400; username conflict returned; no additional account was stored.",
        "test": "test_register_duplicate_username_returns_400_without_another_account"
    },
    {
        "id": "AM_REG_04",
        "group": "Register",
        "description": "Reject an invalid email address",
        "procedure": "1. POST /api/auth/register with email=\"invalid\" and otherwise valid fields.\n2. Inspect the response and database.",
        "expected": "HTTP 422 with an email validation error; no account stored.",
        "actual": "HTTP 422 for the invalid email; account count remained zero.",
        "test": "test_register_rejects_invalid_input_before_database_writes"
    },
    {
        "id": "AM_REG_05",
        "group": "Register",
        "description": "Reject a username shorter than three characters",
        "procedure": "1. POST /api/auth/register with username=\"ab\" and valid email/password.\n2. Inspect the response and database.",
        "expected": "HTTP 422; username length validation fails; no account stored.",
        "actual": "HTTP 422 for username=\"ab\"; account count remained zero.",
        "test": "test_register_rejects_invalid_input_before_database_writes"
    },
    {
        "id": "AM_REG_06",
        "group": "Register",
        "description": "Reject spaces in a username",
        "procedure": "1. POST /api/auth/register with username=\"bad username\".\n2. Keep email/password valid.\n3. Inspect the response.",
        "expected": "HTTP 422; username character validation fails; no account stored.",
        "actual": "HTTP 422 for a username containing a space; account count remained zero.",
        "test": "test_register_rejects_invalid_input_before_database_writes"
    },
    {
        "id": "AM_REG_07",
        "group": "Register",
        "description": "Reject a password shorter than eight characters",
        "procedure": "1. POST /api/auth/register with password=\"short\".\n2. Keep email/username valid.\n3. Inspect the response and database.",
        "expected": "HTTP 422; password minimum-length validation fails; no account stored.",
        "actual": "HTTP 422 for the short password; account count remained zero.",
        "test": "test_register_rejects_invalid_input_before_database_writes"
    },
    {
        "id": "AM_REG_08",
        "group": "Register",
        "description": "Reject missing registration fields",
        "procedure": "1. POST /api/auth/register, omitting email, username and password one at a time.\n2. Also submit an empty JSON object.\n3. Inspect all responses and account count.",
        "expected": "Every request returns HTTP 422; no account stored.",
        "actual": "All missing-field and empty-body requests returned HTTP 422; no account was stored.",
        "test": "test_register_missing_required_fields_returns_422"
    },
    {
        "id": "AM_REG_09",
        "group": "Register",
        "description": "Accept username lengths of exactly three and fifty characters",
        "procedure": "1. Use separate new emails.\n2. Register username=\"abc\".\n3. Register another username consisting of fifty \"a\" characters.",
        "expected": "Both requests return HTTP 201; usernames are preserved; two accounts stored.",
        "actual": "Both boundary-length usernames were accepted with HTTP 201 and preserved; two accounts were stored.",
        "test": "test_register_username_at_both_length_boundaries"
    },
    {
        "id": "AM_REG_10",
        "group": "Register",
        "description": "Reject a username longer than fifty characters",
        "procedure": "1. POST /api/auth/register with a username of fifty-one \"a\" characters.\n2. Keep other fields valid.\n3. Inspect account count.",
        "expected": "HTTP 422; maximum username length exceeded; no account stored.",
        "actual": "HTTP 422 for the fifty-one-character username; account count remained zero.",
        "test": "test_register_rejects_username_longer_than_50_characters"
    },
    {
        "id": "AM_REG_11",
        "group": "Register",
        "description": "Accept and authenticate an eight-character password",
        "procedure": "1. Register a new account with password=\"abcdefgh\".\n2. POST /api/auth/login with the same email and password.",
        "expected": "Register HTTP 201; login HTTP 200. The contract allows nonblank passwords of 8–1024 characters without requiring mixed character classes.",
        "actual": "The eight-character password was accepted; registration returned HTTP 201 and login returned HTTP 200.",
        "test": "test_register_accepts_eight_character_password_and_can_login"
    },
    {
        "id": "AM_REG_12",
        "group": "Register",
        "description": "Prevent client-controlled administrator privileges",
        "procedure": "1. Register valid new credentials with additional role_id=2, role=\"admin\" and is_active=false fields.\n2. Inspect the stored account.",
        "expected": "HTTP 201; ignored extra fields cannot change server-assigned role_id=1 and is_active=true.",
        "actual": "HTTP 201; the stored account retained role_id=1 and is_active=true.",
        "test": "test_register_cannot_assign_a_privileged_role_from_client_input"
    },
    {
        "id": "AM_REG_13",
        "group": "Register",
        "description": "Keep registration independent of optional SMTP delivery",
        "procedure": "1. In the isolated test, make the password-reset email function raise an SMTP error if called.\n2. Submit a valid registration.\n3. Check the response and email-function calls.",
        "expected": "HTTP 201; registration does not invoke password-reset email delivery.",
        "actual": "HTTP 201; the optional email function was not called.",
        "test": "test_registration_does_not_depend_on_optional_email_delivery"
    },
    {
        "id": "AM_REG_14",
        "group": "Register",
        "description": "Reject null registration fields",
        "procedure": "1. POST /api/auth/register with email=null, username=null and password=null in separate requests.\n2. Keep remaining fields valid.\n3. Check account count.",
        "expected": "Each request returns HTTP 422; no account stored.",
        "actual": "All null-field requests returned HTTP 422; account count remained zero.",
        "test": "test_register_null_required_fields_returns_422"
    },
    {
        "id": "AM_REG_15",
        "group": "Register",
        "description": "Reject empty registration fields",
        "procedure": "1. POST /api/auth/register with empty email, username and password strings in separate requests.\n2. Keep remaining fields valid.\n3. Check account count.",
        "expected": "Each request returns HTTP 422; no account stored.",
        "actual": "All empty-field requests returned HTTP 422; account count remained zero.",
        "test": "test_register_empty_required_fields_returns_422"
    },
    {
        "id": "AM_REG_16",
        "group": "Register",
        "description": "Normalize email and reject case-variant duplicates",
        "procedure": "1. Register email=\"  Mixed.Case@EXAMPLE.COM  \".\n2. Retry with email=\"MIXED.CASE@example.com\" and a different username.",
        "expected": "First request HTTP 201 with email=\"mixed.case@example.com\"; duplicate request HTTP 400; one account stored.",
        "actual": "Email was normalized to lowercase; the case-variant duplicate returned HTTP 400; one account was stored.",
        "test": "test_register_normalizes_email_and_rejects_case_variant_duplicates"
    },
    {
        "id": "AM_REG_17",
        "group": "Register",
        "description": "Handle a competing duplicate registration safely",
        "procedure": "1. In the isolated test, simulate a competing registration winning after availability checks.\n2. Allow the actual database uniqueness constraint to reject the duplicate insert.\n3. Inspect the response and account count.",
        "expected": "HTTP 400 with the email conflict; failed transaction rolled back; only the original account remains.",
        "actual": "The database uniqueness conflict was rolled back and translated to HTTP 400; only one account remained.",
        "test": "test_register_recovers_a_duplicate_that_wins_after_availability_check"
    },
    {
        "id": "AM_REG_18",
        "group": "Register",
        "description": "Roll back a failed registration commit",
        "procedure": "1. Make the isolated session commit raise a database error after flush.\n2. Submit valid registration data.\n3. Inspect the response and account count.",
        "expected": "HTTP 500 with a safe registration error; no partial account remains; private database details are absent.",
        "actual": "HTTP 500; no private database error appeared in the response; the flushed account was rolled back.",
        "test": "test_register_rolls_back_failed_commit_without_leaking_database_errors"
    },
    {
        "id": "AM_REG_19",
        "group": "Register",
        "description": "Bound password size for registration and login",
        "procedure": "1. POST /api/auth/register with a 1025-character password.\n2. POST /api/auth/login with a 1025-character password.\n3. Inspect the responses and database.",
        "expected": "Both requests return HTTP 422; no account stored.",
        "actual": "Both oversized-password requests returned HTTP 422; account count remained zero.",
        "test": "test_register_and_login_reject_excessively_large_passwords"
    },
    {
        "id": "AM_REG_20",
        "group": "Register",
        "description": "Reject a password containing only whitespace",
        "procedure": "1. POST /api/auth/register with a password containing eight spaces.\n2. Keep email/username valid.\n3. Inspect the response and database.",
        "expected": "HTTP 422; a password must contain a non-whitespace character; no account stored.",
        "actual": "HTTP 422; the all-whitespace password was rejected and no account was stored.",
        "test": "test_register_rejects_an_all_whitespace_password"
    },
    {
        "id": "AM_REG_21",
        "group": "Register",
        "description": "Allocate distinct account IDs through the database",
        "procedure": "1. Register two accounts with different email/username pairs.\n2. Compare returned account IDs and account count.",
        "expected": "Both requests return HTTP 201; IDs differ; two accounts stored. ID allocation uses the database allocator.",
        "actual": "Both registrations returned HTTP 201; IDs were distinct and two accounts were stored.",
        "test": "test_register_allocates_distinct_ids_without_application_max_id_query"
    },
    {
        "id": "AM_LOG_01",
        "group": "Login",
        "description": "Log in by email and authenticate the issued JWT",
        "procedure": "1. Register a valid account.\n2. POST /api/auth/login with its email/password.\n3. Call GET /api/auth/me using Authorization: Bearer <access_token>.",
        "expected": "Login HTTP 200; token_type=\"bearer\", user.role=\"user\"; authenticated /me HTTP 200 with matching account_id/email.",
        "actual": "Email login and the protected endpoint both returned HTTP 200; the token and account identity were verified.",
        "test": "test_registered_user_can_login_and_access_the_protected_auth_endpoint"
    },
    {
        "id": "AM_LOG_02",
        "group": "Login",
        "description": "Reject an incorrect password",
        "procedure": "1. Register an account.\n2. POST /api/auth/login with the correct email and password=\"WrongPassword@2026\".",
        "expected": "HTTP 401; detail=\"Incorrect email or password.\"; no access token issued.",
        "actual": "The incorrect password was rejected with HTTP 401.",
        "test": "test_login_with_the_wrong_password_returns_401"
    },
    {
        "id": "AM_LOG_03",
        "group": "Login",
        "description": "Reject an inactive account",
        "procedure": "1. Register a test-only account.\n2. Set is_active=false for that account in the isolated database.\n3. Log in with the correct credentials.",
        "expected": "HTTP 403; detail=\"This account is locked or not activated.\"; no access token issued.",
        "actual": "The inactive account was rejected with HTTP 403 despite correct credentials.",
        "test": "test_inactive_account_cannot_login"
    },
    {
        "id": "AM_LOG_04",
        "group": "Login",
        "description": "Reject an unknown email without disclosing account existence",
        "procedure": "1. POST /api/auth/login with a valid but unregistered email and any password.\n2. Inspect the response.",
        "expected": "HTTP 401; detail=\"Incorrect email or password.\"; no access token issued.",
        "actual": "HTTP 401 with the generic credentials error; no access token was returned.",
        "test": "test_login_nonexistent_email_returns_the_same_generic_401"
    },
    {
        "id": "AM_LOG_05",
        "group": "Login",
        "description": "Reject missing login credentials",
        "procedure": "1. POST /api/auth/login with an empty body.\n2. Send email without password.\n3. Send password without an identifier.",
        "expected": "Every request returns HTTP 422.",
        "actual": "All missing-credential requests returned HTTP 422.",
        "test": "test_login_missing_required_fields_returns_422"
    },
    {
        "id": "AM_LOG_06",
        "group": "Login",
        "description": "Reject an empty, invalid or null email without another identifier",
        "procedure": "1. POST /api/auth/login with email=\"\", email=\"not-an-email\" and email=null in separate requests.\n2. Supply a password, with no username or username_or_email.",
        "expected": "Every request returns HTTP 422; no access token issued.",
        "actual": "All invalid-email or missing-identifier requests returned HTTP 422.",
        "test": "test_login_rejects_empty_or_invalid_email"
    },
    {
        "id": "AM_LOG_07",
        "group": "Login",
        "description": "Log in using a username",
        "procedure": "1. Register a valid account.\n2. POST /api/auth/login with its username and password, omitting email.\n3. Verify the returned JWT.",
        "expected": "HTTP 200; the verified token subject matches the registered account ID.",
        "actual": "Username login returned HTTP 200; the verified JWT subject matched the account ID.",
        "test": "test_login_with_username_issues_a_working_jwt"
    },
    {
        "id": "AM_LOG_08",
        "group": "Login",
        "description": "Normalize email case and surrounding whitespace during login",
        "procedure": "1. Register registration@example.com.\n2. POST /api/auth/login with email=\"  REGISTRATION@EXAMPLE.COM  \" and the correct password.",
        "expected": "HTTP 200; the response identifies the registered lowercase email account.",
        "actual": "HTTP 200; normalization resolved the same registered account.",
        "test": "test_login_normalizes_mixed_case_email"
    },
    {
        "id": "AM_LOG_09",
        "group": "Login",
        "description": "Treat passwords as case-sensitive",
        "procedure": "1. Register with a mixed-case password.\n2. Log in with the correct email and a lowercase version of that password.",
        "expected": "HTTP 401; no access token issued.",
        "actual": "The case-changed password returned HTTP 401; no token was returned.",
        "test": "test_login_password_is_case_sensitive"
    },
    {
        "id": "AM_LOG_10",
        "group": "Login",
        "description": "Keep password material out of login responses and JWT claims",
        "procedure": "1. Log in with valid credentials.\n2. Inspect the response and verify/decode the JWT locally.",
        "expected": "HTTP 200; response contains access_token, token_type and user(account_id,email,username,role). JWT has expiry and the expected role. Neither response nor claims contains password/password_hash.",
        "actual": "The response shape, expiry and role were verified; password and hash fields were absent from the response and JWT.",
        "test": "test_login_response_and_jwt_never_expose_password_or_hash"
    },
    {
        "id": "AM_LOG_11",
        "group": "Login",
        "description": "Support the compatibility username_or_email field",
        "procedure": "1. Register an account.\n2. Log in with username_or_email set to its uppercase email.\n3. Repeat with username_or_email set to its username.",
        "expected": "Both requests return HTTP 200.",
        "actual": "Both email and username compatibility requests returned HTTP 200.",
        "test": "test_login_legacy_username_or_email_supports_both_identifiers"
    },
    {
        "id": "AM_LOG_12",
        "group": "Login",
        "description": "Reject ambiguous login identifiers",
        "procedure": "1. POST /api/auth/login with both email and username.\n2. Repeat with email and username_or_email.\n3. Supply a password in both requests.",
        "expected": "Both requests return HTTP 422; exactly one identifier is accepted.",
        "actual": "Both ambiguous-identifier requests returned HTTP 422.",
        "test": "test_login_rejects_ambiguous_identifiers"
    },
    {
        "id": "AM_LOG_13",
        "group": "Login",
        "description": "Reject a temporary account lock that has not expired",
        "procedure": "1. Register a test account.\n2. Set locked_until to a future timestamp in the test database.\n3. Log in with correct credentials.",
        "expected": "HTTP 403; detail=\"This account is locked or not activated.\"; no access token issued.",
        "actual": "The future temporary lock returned HTTP 403 with the expected account-unavailable detail.",
        "test": "test_login_rejects_temporarily_locked_account"
    },
    {
        "id": "AM_LOG_14",
        "group": "Login",
        "description": "Allow login after a temporary lock expires",
        "procedure": "1. Register a test account.\n2. Set locked_until to a past timestamp.\n3. Log in with correct credentials.",
        "expected": "HTTP 200. Lock expiry is evaluated against the database clock and timezone.",
        "actual": "Login returned HTTP 200 after the temporary lock expired.",
        "test": "test_login_accepts_an_expired_temporary_lock"
    },
    {
        "id": "AM_LOG_15",
        "group": "Login",
        "description": "Handle malformed stored password hashes safely",
        "procedure": "1. Register a test account.\n2. Replace only its stored hash with an invalid legacy value in the isolated database.\n3. Submit a login request.",
        "expected": "HTTP 401 with the generic credentials error; no unhandled hash exception or access token.",
        "actual": "The malformed stored hash returned HTTP 401 with the generic error instead of HTTP 500.",
        "test": "test_login_invalid_stored_hash_returns_401_instead_of_crashing"
    },
    {
        "id": "AM_LOG_16",
        "group": "Login",
        "description": "Reject a password-reset token as an access token",
        "procedure": "1. Generate a password-reset JWT in the isolated test.\n2. Call GET /api/auth/me using that token as a bearer credential.",
        "expected": "HTTP 401; a reset token cannot authenticate a protected endpoint.",
        "actual": "The password-reset token was rejected by the protected endpoint with HTTP 401.",
        "test": "test_password_reset_token_cannot_authenticate_as_an_access_token"
    }
]

BLACK, WHITE, ORANGE = "FF000000", "FFFFFFFF", "FFED7D31"
THIN = Side(style="thin", color=BLACK)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RESULT_COLORS = {
    "Pass": ("FF107C41", WHITE),
    "Fail": ("FFC00000", WHITE),
    "Not Run": ("FFFFF2CC", "FF7F6000"),
}


def read_evidence(path: Path | None):
    if path is None:
        return {}, "", "No execution evidence supplied."
    raw = Path(path).read_bytes()
    root = ElementTree.fromstring(raw)
    outcomes = defaultdict(list)
    for node in root.iter("testcase"):
        if node.get("classname") != "apps.backend.tests.test_auth_api.AuthAPITests":
            continue
        method = node.get("name", "").split("[", 1)[0].strip()
        if node.find("failure") is not None or node.find("error") is not None:
            outcome = "Fail"
        elif node.find("skipped") is not None:
            outcome = "Not Run"
        else:
            outcome = "Pass"
        outcomes[method].append(outcome)
    if not outcomes:
        raise ValueError("JUnit does not contain the existing AuthAPITests suite.")
    states = {
        method: "Fail" if "Fail" in items else "Not Run" if "Not Run" in items else "Pass"
        for method, items in outcomes.items()
    }
    suite = next(root.iter("testsuite"))
    executed_at = datetime.fromisoformat(suite.attrib["timestamp"])
    if executed_at.tzinfo is None:
        executed_at = executed_at.replace(tzinfo=timezone.utc)
    date = executed_at.astimezone(timezone(timedelta(hours=7))).strftime("%d/%m/%Y")
    evidence = f"JUnit file: {Path(path).name}; SHA256: {hashlib.sha256(raw).hexdigest()}."
    return states, date, evidence


def estimated_height(values, widths, minimum=52):
    longest = 1
    for value, width in zip(values, widths):
        lines = sum(
            max(1, len(wrap(line, width=int(width) - 3)))
            for line in str(value).split("\n")
        )
        longest = max(longest, lines)
    return max(minimum, longest * 16 + 12)


def font(bold=False, color=BLACK, size=11):
    return Font(name="Times New Roman", size=size, bold=bold, color=color)


def merged(sheet, row, start, end, value, bold=False, color=BLACK, centered=False):
    sheet.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    sheet.cell(row, start, value)
    for col in range(start, end + 1):
        cell = sheet.cell(row, col)
        cell.font = font(bold=bold, color=color)
        cell.alignment = Alignment(
            horizontal="center" if centered else "left", vertical="center", wrap_text=True,
        )
        cell.fill = PatternFill("solid", fgColor=WHITE)
        cell.border = BORDER


def build_report(output=OUTPUT, junit=None):
    if len(CASES) != len({case["id"] for case in CASES}):
        raise ValueError("FUT IDs must be unique.")
    states, date, evidence = read_evidence(junit)
    counts = Counter(states.get(case["test"], "Not Run") for case in CASES)
    book = Workbook()
    sheet = book.active
    sheet.title = "Account_Management"
    book.properties.creator = TESTER
    book.properties.lastModifiedBy = TESTER
    book.properties.title = "FUT - Backend Login and Register"
    book.properties.subject = "Account Management: FastAPI Register and Login"
    book.properties.description = (
        f"Backend source: {SOURCE_REVISION}. {evidence} "
        "Case status is derived from executed pytest evidence."
    )
    for col, width in enumerate(WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width

    metadata = (
        ("Module Code", "Account Management - Backend Register & Login"),
        ("Tester", TESTER),
        (
            "Pre-condition",
            "Manual API execution: FastAPI at http://127.0.0.1:8000; migrated heritagegraph "
            "database; account and roles tables; USER role id=1; dedicated test accounts.\n"
            "Recorded automated evidence: isolated SQLite database, real ORM transactions, "
            "Argon2 hashing and JWT validation; existing tests/test_auth_api.py under apps/backend.",
        ),
    )
    for row, (label, value) in enumerate(metadata, 1):
        cell = sheet.cell(row, 1, label)
        cell.font = font(bold=True)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        merged(sheet, row, 2, 7, value, centered=row < 3)
        sheet.row_dimensions[row].height = 26 if row < 3 else 66

    sheet["A1"].comment = Comment(
        f"Scope: POST /api/auth/register and POST /api/auth/login. Backend source {SOURCE_REVISION}. "
        "GET /api/auth/me validates the token issued by login. "
        "Each result references an existing automated test method. "
        "Multiple FUT cases may share a method containing distinct subtests.",
        TESTER,
    )
    for address, value in (("A4", "Pass"), ("B4", "Fail")):
        cell = sheet[address]
        cell.value = value
        cell.font = font(bold=True, size=12)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    merged(sheet, 4, 3, 7, "Number of test case", bold=True, centered=True)
    for address, value in (("A5", counts["Pass"]), ("B5", counts["Fail"])):
        cell = sheet[address]
        cell.value = value
        cell.font = font()
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    merged(sheet, 5, 3, 7, len(CASES), centered=True)
    sheet["C5"].comment = Comment(
        f"{counts['Pass']} Pass + {counts['Fail']} Fail + {counts['Not Run']} Not Run = {len(CASES)}. "
        "Counts reflect supplied JUnit evidence at generation time. Regenerate after another run.",
        TESTER,
    )
    for row in (4, 5):
        sheet.row_dimensions[row].height = 25
    sheet.row_dimensions[6].height = 12
    for col, title in enumerate(HEADERS, 1):
        cell = sheet.cell(7, col, title)
        cell.font = font(bold=True, size=12)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[7].height = 36
    sheet["F7"].comment = Comment(
        "Pass means the mapped test/subtests completed successfully, including expected negative responses. "
        "A skipped or missing test is Not Run. The workbook generator does not execute tests. " + evidence,
        TESTER,
    )
    dropdown = DataValidation(type="list", formula1='"Pass,Fail,Not Run"', allow_blank=False)
    dropdown.errorTitle = "Invalid result"
    dropdown.error = "Choose Pass, Fail or Not Run."
    dropdown.showErrorMessage = True
    dropdown.errorStyle = "stop"
    sheet.add_data_validation(dropdown)

    row, group = 8, None
    for case in CASES:
        if case["group"] != group:
            group = case["group"]
            merged(sheet, row, 1, 7, group, bold=True, color=ORANGE, centered=True)
            sheet.row_dimensions[row].height = 24
            row += 1
        result = states.get(case["test"], "Not Run")
        actual = case["actual"] if result == "Pass" else (
            "The mapped automated test failed. Review its JUnit failure details."
            if result == "Fail" else "Not executed or skipped; record actual results after execution."
        )
        values = [
            case["id"], case["description"], case["procedure"], case["expected"],
            actual, result, date if result != "Not Run" else "",
        ]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row, col)
            cell.font = font(bold=col in (1, 6))
            cell.border = BORDER
            cell.alignment = Alignment(
                horizontal="center" if col in (1, 6, 7) else "left",
                vertical="center" if col in (1, 6, 7) else "top", wrap_text=True,
            )
            cell.value = datetime.strptime(value, "%d/%m/%Y") if col == 7 and value else value or None
            if col == 7:
                cell.number_format = "dd/mm/yyyy"
            if col == 6:
                fill, text = RESULT_COLORS[result]
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = font(bold=True, color=text)
                dropdown.add(cell)
                for state, (bg, fg) in RESULT_COLORS.items():
                    sheet.conditional_formatting.add(
                        cell.coordinate, CellIsRule(
                            operator="equal", formula=[f'"{state}"'],
                            fill=PatternFill("solid", fgColor=bg), font=font(bold=True, color=fg),
                        ),
                    )
        sheet.cell(row, 5).comment = Comment(
            f"apps/backend/tests/test_auth_api.py::AuthAPITests::{case['test']}\n"
            f"Execution date (Asia/Saigon): {date or 'Not Run'}. {evidence}\n"
            "Automated evidence uses isolated SQLite with real ORM, Argon2 and JWTs. "
            "Separate localhost/PostgreSQL checks also validated concurrent registration and lock expiry.",
            TESTER,
        )
        sheet.row_dimensions[row].height = estimated_height(values, WIDTHS)
        row += 1

    sheet.auto_filter.ref = f"A7:G{row - 1}"
    notes = (
        f"Summary: {len(CASES)} FUT cases | {counts['Pass']} Pass | {counts['Fail']} Fail | "
        f"{counts['Not Run']} Not Run. Counts refer to FUT cases, not distinct test methods.",
        "Results are derived from the supplied pytest JUnit report. The existing auth suite includes "
        "real database writes, Argon2 and JWT verification; external email delivery is not required.",
        "Backend contract: login accepts one email/username/username_or_email; registration creates an active "
        "USER account. Email verification and frontend navigation are outside this FUT scope.",
    )
    row += 1
    for note in notes:
        merged(sheet, row, 1, 7, note)
        sheet.row_dimensions[row].height = 36
        row += 1
    sheet.freeze_panes = "B8"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 75
    sheet.print_area = f"A1:G{row - 1}"
    sheet.print_title_rows = "7:7"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4)
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.left.text = TESTER
    sheet.oddFooter.right.text = "Page &P / &N"
    book.save(output)
    book.close()
    return counts


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--junit", type=Path, help="pytest JUnit XML containing AuthAPITests")
    parser.add_argument("--output", type=Path, default=OUTPUT, help="Destination workbook")
    args = parser.parse_args()
    counts = build_report(args.output, args.junit)
    print(f"Created: {args.output}")
    print(f"Pass: {counts['Pass']} | Fail: {counts['Fail']} | Not Run: {counts['Not Run']} | Total: {len(CASES)}")
    return 1 if counts["Fail"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
